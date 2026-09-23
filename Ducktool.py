# ============================================================
# 肉鸭采食行为 + 体重 + 生产性能综合分析工具  Python V5.0
# ============================================================

import io
import re
import math
import warnings
from datetime import datetime, timedelta

import numpy as np
import pandas as pd
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
from scipy import stats
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Alignment

warnings.filterwarnings("ignore")

st.set_page_config(
    page_title="肉鸭采食行为与生产性能综合分析工具",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ============================================================
# 调色板
# ============================================================
PALETTES = {
    "单色-浅蓝": ["#9ecae1"],
    "单色-深蓝": ["#2c7fb8"],
    "单色-绿":   ["#31a354"],
    "单色-橙":   ["#e6550d"],
    "单色-紫":   ["#756bb1"],
    "单色-灰":   ["#636363"],
    "暖色系":    ["#a30543", "#f36f43", "#fbda83"],
    "冷色系":    ["#e9f4a3", "#80cba4", "#4965b0"],
    "六色混合":  ["#a30543", "#f36f43", "#fbda83", "#e9f4a3", "#80cba4", "#4965b0"],
    "蓝色渐变":  ["#c6dbef", "#9ecae1", "#6baed6", "#4292c6", "#2171b5", "#08519c"],
    "黄绿渐变":  ["#ffffcc", "#c2e699", "#78c679", "#31a354", "#006837"],
    "红紫蓝绿":  ["#e41a1c", "#377eb8", "#4daf4a", "#984ea3", "#ff7f00", "#ffff33"],
    "Pastel":    ["#fbb4ae", "#b3cde3", "#ccebc5", "#decbe4", "#fed9a6", "#ffffcc"],
    "深色系":    ["#1b9e77", "#d95f02", "#7570b3", "#e7298a", "#66a61e", "#e6ab02"],
}

HEAT_PALETTES = {
    "红-白-蓝(默认)": ["#b2182b", "#f7f7f7", "#2166ac"],
    "红-白-绿":       ["#d73027", "#ffffbf", "#1a9850"],
    "紫-白-橙":       ["#762a83", "#f7f7f7", "#e08214"],
    "蓝-白-红":       ["#2166ac", "#f7f7f7", "#b2182b"],
    "深红-白-深蓝":   ["#67001f", "#f7f7f7", "#053061"],
    "橙-白-紫":       ["#f1a340", "#f7f7f7", "#998ec3"],
    "绿-白-红":       ["#1a9850", "#ffffbf", "#d73027"],
}


def seq_colors(pal_name, n):
    p = PALETTES.get(pal_name, ["#2c7fb8"])
    return [p[i % len(p)] for i in range(n)]


def valid_hex(x, fallback="#9ecae1"):
    if x is None:
        return fallback
    x = str(x).strip()
    if re.match(r"^#[0-9A-Fa-f]{6}$", x):
        return x
    if re.match(r"^[0-9A-Fa-f]{6}$", x):
        return "#" + x
    return fallback


def format_p(x):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    if x < 0.001:
        return "<0.001"
    return f"{x:.3f}"


# ============================================================
# 文件读取 / 字段识别
# ============================================================
FEED_PATTERNS = {
    "animal": [
        r"^耳标$", r"耳标", r"耳号", r"animal[\s_\-]*id", r"animal[\s_\-]*no",
        r"individual", r"tag", r"^id$", r"编号",
    ],
    "time1": [
        r"^记录时间1$", r"记录时间\s*1", r"记录.*时间.*1", r"time\s*1",
        r"start[\s_\-]*time", r"开始时间",
    ],
    "time2": [
        r"^记录时间2$", r"记录时间\s*2", r"记录.*时间.*2", r"time\s*2",
        r"end[\s_\-]*time", r"结束时间",
    ],
    "intake": [
        r"^处理后采食量\s*[\(（]?\s*[gG]\s*[\)）]?$",
        r"处理后.*采食量", r"processed.*feed.*intake",
        r"^采食量", r"feed[\s_\-]*intake",
        r"食槽重.*差", r"采食.*重",
    ],
    "duration": [
        r"^采食时长$", r"采食.*时长", r"feeding.*duration",
        r"^duration$", r"持续.*时间", r"用时",
    ],
}

WEIGHT_PATTERNS = {
    "animal": [
        r"^耳标$", r"耳标", r"耳号", r"animal[\s_\-]*id", r"animal[\s_\-]*no",
        r"individual", r"tag", r"^id$", r"编号",
    ],
    "time1": [
        r"^记录时间1$", r"记录时间\s*1", r"记录.*时间.*1", r"time\s*1",
        r"start[\s_\-]*time", r"开始时间",
    ],
    "time2": [
        r"^记录时间2$", r"记录时间\s*2", r"记录.*时间.*2", r"time\s*2",
        r"end[\s_\-]*time", r"结束时间",
    ],
    "bw1": [
        r"^体重\s*1\s*[\(（]?\s*[kK][gG]\s*[\)）]?$",
        r"^体重\s*1$", r"weight\s*1", r"^初重$", r"^初始体重",
    ],
    "bw2": [
        r"^体重\s*2\s*[\(（]?\s*[kK][gG]\s*[\)）]?$",
        r"^体重\s*2$", r"weight\s*2", r"^中重$",
    ],
    "bw3": [
        r"^体重\s*3\s*[\(（]?\s*[kK][gG]\s*[\)）]?$",
        r"^体重\s*3$", r"weight\s*3", r"^末重$", r"^终重$",
    ],
    "mean_bw": [
        r"^平均体重\s*[\(（]?\s*[kK][gG]\s*[\)）]?$",
        r"^平均体重", r"mean[\s_\-]*weight", r"average[\s_\-]*weight",
        r"^体重$",
    ],
}

FEED_REQUIRED = ["animal", "time1"]
WEIGHT_REQUIRED = ["animal", "time1"]


def clean_names(nms):
    """列名清洗：全角转半角、统一括号、去空白、去换行。"""
    def _clean_one(n):
        s = str(n)
        s = s.replace("\u00a0", " ").replace("\u3000", " ")
        s = s.replace("（", "(").replace("）", ")")
        s = s.replace("ｇ", "g").replace("Ｇ", "G")
        s = s.replace("ｋ", "k").replace("Ｋ", "K")
        s = s.replace("\r", "").replace("\n", "").replace("\t", "")
        s = re.sub(r"\s+", " ", s).strip()
        return s
    return [_clean_one(n) for n in nms]


def find_col(df, patterns, required=False):
    nms = list(df.columns)
    for p in patterns:
        for n in nms:
            if re.search(p, str(n).strip(), flags=re.IGNORECASE):
                return n
    if required:
        raise ValueError(
            f"无法找到需要的字段。\n候选关键词：{' / '.join(patterns)}\n"
            f"当前字段：{' | '.join(map(str, nms))}"
        )
    return None


def sheet_matches_dataset(header_df, dataset_type, return_detail=False):
    if dataset_type == "feed":
        pats, required = FEED_PATTERNS, FEED_REQUIRED
    else:
        pats, required = WEIGHT_PATTERNS, WEIGHT_REQUIRED

    cols = clean_names(header_df.columns)
    tmp = pd.DataFrame(columns=cols)
    detail = {}
    for key, p in pats.items():
        found = find_col(tmp, p, required=False)
        detail[key] = found

    matched = all(detail.get(k) is not None for k in required)
    if return_detail:
        return matched, detail
    return matched


def detect_header_row(path_or_buf, sheet_name, dataset_type, max_scan=6):
    pats = FEED_PATTERNS if dataset_type == "feed" else WEIGHT_PATTERNS
    try:
        raw = pd.read_excel(path_or_buf, sheet_name=sheet_name,
                            header=None, nrows=max_scan, dtype=object)
    except Exception:
        return 0
    if len(raw) == 0:
        return 0
    best_row, best_score = 0, -1
    for i in range(len(raw)):
        cols = clean_names(raw.iloc[i].tolist())
        tmp = pd.DataFrame(columns=cols)
        score = sum(1 for p in pats.values() if find_col(tmp, p, required=False) is not None)
        if score > best_score:
            best_score, best_row = score, i
    return best_row


def read_excel_clean(path_or_buf, sheet_name=0, dataset_type=None):
    header_row = 0
    if dataset_type:
        try:
            header_row = detect_header_row(path_or_buf, sheet_name, dataset_type)
        except Exception:
            header_row = 0
    df = pd.read_excel(path_or_buf, sheet_name=sheet_name,
                       header=header_row, dtype=object)
    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = clean_names(df.columns)
    return df.reset_index(drop=True)


def read_matching_excel_files(files, dataset_type):
    """files: list of (name, bytes)。能读到多少字段就读多少，必需字段不全的 sheet 跳过并记录原因。"""
    if not files:
        raise ValueError("请至少上传1个采食 Excel 文件。" if dataset_type == "feed"
                         else "请至少上传1个体重 Excel 文件。")

    dataset_label = "采食数据" if dataset_type == "feed" else "体重数据"
    data_list, qc_list = [], []

    for name, raw in files:
        try:
            xls = pd.ExcelFile(io.BytesIO(raw))
        except Exception as e:
            qc_list.append({"数据集": dataset_label, "Source_File": name, "Source_Sheet": "-",
                            "状态": "跳过：文件读取失败", "读取记录数": 0, "说明": str(e)})
            continue

        for sheet in xls.sheet_names:
            try:
                header = pd.read_excel(xls, sheet_name=sheet, nrows=6)
                header.columns = clean_names(header.columns)
            except Exception as e:
                qc_list.append({"数据集": dataset_label, "Source_File": name, "Source_Sheet": sheet,
                                "状态": "跳过：工作表读取失败", "读取记录数": 0, "说明": str(e)})
                continue

            matched, detail = sheet_matches_dataset(header, dataset_type, return_detail=True)
            if not matched:
                missing = [k for k, v in detail.items() if v is None]
                qc_list.append({
                    "数据集": dataset_label, "Source_File": name, "Source_Sheet": sheet,
                    "状态": "跳过：必需字段缺失", "读取记录数": 0,
                    "说明": f"缺失字段: {missing}；候选列: {list(header.columns)[:20]}",
                })
                continue

            try:
                tmp = read_excel_clean(xls, sheet, dataset_type=dataset_type)
            except Exception as e:
                qc_list.append({"数据集": dataset_label, "Source_File": name, "Source_Sheet": sheet,
                                "状态": "跳过：工作表读取失败", "读取记录数": 0, "说明": str(e)})
                continue

            if len(tmp) == 0:
                qc_list.append({"数据集": dataset_label, "Source_File": name, "Source_Sheet": sheet,
                                "状态": "跳过：工作表为空", "读取记录数": 0,
                                "说明": "字段匹配，但没有有效记录"})
                continue

            tmp["_Source_File"] = name
            tmp["_Source_Sheet"] = sheet
            data_list.append(tmp)

            matched_ok = [k for k, v in detail.items() if v is not None]
            matched_ng = [k for k, v in detail.items() if v is None]
            qc_list.append({
                "数据集": dataset_label, "Source_File": name, "Source_Sheet": sheet,
                "状态": "已读取", "读取记录数": len(tmp),
                "说明": f"匹配字段: {matched_ok}" + (f" | 缺失(将置空): {matched_ng}" if matched_ng else ""),
            })

    qc_df = pd.DataFrame(qc_list)
    if not data_list:
        return pd.DataFrame(), qc_df
    return pd.concat(data_list, ignore_index=True), qc_df


# ============================================================
# 数值 / 时间 / 时长解析
# ============================================================
def safe_numeric(x):
    if isinstance(x, pd.Series):
        s = x.astype(str).str.replace(",", "", regex=False).str.strip()
        s = s.replace({"": None, "NA": None, "NaN": None, "NULL": None, "null": None, "-": None})
        return pd.to_numeric(s, errors="coerce")
    return x


def parse_datetime_flexible(x):
    if isinstance(x, pd.Series):
        s = x.copy()
        if pd.api.types.is_numeric_dtype(s):
            return pd.to_datetime(s, unit="D", origin="1899-12-30", errors="coerce")
        s = s.astype(str).str.strip().replace({"": None, "NA": None, "NaN": None, "NULL": None, "null": None})
        out = pd.to_datetime(s, errors="coerce")
        if out.isna().any():
            formats = [
                "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y.%m.%d %H:%M:%S",
                "%Y/%m/%d %H:%M", "%Y-%m-%d %H:%M", "%Y.%m.%d %H:%M",
                "%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d",
                "%m/%d/%Y %H:%M:%S", "%m-%d-%Y %H:%M:%S",
                "%m/%d/%Y %H:%M", "%m-%d-%Y %H:%M",
                "%m/%d/%Y", "%m-%d-%Y",
                "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S",
                "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M",
                "%d/%m/%Y", "%d-%m-%Y",
            ]
            for f in formats:
                mask = out.isna() & s.notna()
                if not mask.any():
                    break
                try:
                    out.loc[mask] = pd.to_datetime(s[mask], format=f, errors="coerce")
                except Exception:
                    pass
        return out
    return x


def parse_duration_seconds(x):
    """支持 1.09:33:38 / 09:33:38 / 33:33:38 / 09:33 / 数字(天)"""
    if isinstance(x, pd.Series):
        s = x.astype(str).str.strip().replace({"": None, "NA": None, "NaN": None, "NULL": None, "null": None})

        def _one(v):
            if v is None or (isinstance(v, float) and math.isnan(v)):
                return np.nan
            v = str(v).strip()
            m = re.match(r"^(\d+)\.(\d{1,2}):(\d{2}):(\d{2})$", v)
            if m:
                d, h, mi, se = map(int, m.groups())
                return d * 86400 + h * 3600 + mi * 60 + se
            m = re.match(r"^(\d+):(\d{2}):(\d{2})$", v)
            if m:
                h, mi, se = map(int, m.groups())
                return h * 3600 + mi * 60 + se
            m = re.match(r"^(\d+):(\d{2})$", v)
            if m:
                h, mi = map(int, m.groups())
                return h * 3600 + mi * 60
            m = re.match(r"^(\d+(\.\d+)?)$", v)
            if m:
                return float(m.group(1)) * 86400
            return np.nan

        return s.map(_one).astype(float)

    if isinstance(x, (int, float)) and not isinstance(x, bool):
        return float(x) * 86400
    return x


def parse_clock_hhmm(x, default="06:00"):
    if x is None:
        return default
    x = str(x).strip()
    if re.match(r"^\d{1,2}:\d{2}:\d{2}$", x) or re.match(r"^\d{1,2}:\d{2}$", x):
        parts = x.split(":")
        return f"{int(parts[0]):02d}:{int(parts[1]):02d}"
    return default


def clock_to_minutes(x):
    hhmm = parse_clock_hhmm(x)
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


# ============================================================
# 核心计算
# ============================================================
def standardize_feed(df):
    """所有字段宽松匹配；除 Animal_ID / Time1 外都可为 NaN。"""
    def _col(key, required=False):
        return find_col(df, FEED_PATTERNS[key], required=required)

    animal_col   = _col("animal", required=True)
    time1_col    = _col("time1", required=True)
    time2_col    = _col("time2")
    intake_col   = _col("intake")
    duration_col = _col("duration")
    group_col    = find_col(df, [r"^群组$", r"群组", r"group"])
    house_col    = find_col(df, [r"^栋舍$", r"栋舍", r"house"])
    pen_col      = find_col(df, [r"^栏圈$", r"栏圈", r"pen"])

    def _get(col):
        return df[col] if col else pd.Series([np.nan] * len(df), index=df.index)

    out = pd.DataFrame({
        "Animal_ID": df[animal_col].astype(str),
        "Group": _get(group_col).astype(str),
        "House": _get(house_col).astype(str),
        "Pen": _get(pen_col).astype(str),
        "Time1_raw": df[time1_col],
        "Time2_raw": _get(time2_col),
        "Feed_Intake_g": safe_numeric(_get(intake_col)),
        "Duration_raw": _get(duration_col),
        "Source_File": df.get("_Source_File", None),
        "Source_Sheet": df.get("_Source_Sheet", None),
    })
    out["Time1"] = parse_datetime_flexible(out["Time1_raw"])
    out["Time2"] = parse_datetime_flexible(out["Time2_raw"])
    out["Date"] = out["Time1"].dt.date
    out["Feed_Duration_sec"] = parse_duration_seconds(out["Duration_raw"])
    out["Feed_Duration_h"] = out["Feed_Duration_sec"] / 3600
    return out


def standardize_bw(df):
    def _col(key, required=False):
        return find_col(df, WEIGHT_PATTERNS[key], required=required)

    animal_col = _col("animal", required=True)
    time1_col  = _col("time1", required=True)
    time2_col  = _col("time2")
    bw1_col    = _col("bw1")
    bw2_col    = _col("bw2")
    bw3_col    = _col("bw3")
    mean_col   = _col("mean_bw")
    group_col  = find_col(df, [r"^群组$", r"群组", r"group"])
    house_col  = find_col(df, [r"^栋舍$", r"栋舍", r"house"])
    pen_col    = find_col(df, [r"^栏圈$", r"栏圈", r"pen"])

    def _get(col):
        return df[col] if col else pd.Series([np.nan] * len(df), index=df.index)

    out = pd.DataFrame({
        "Animal_ID": df[animal_col].astype(str),
        "Group": _get(group_col).astype(str),
        "House": _get(house_col).astype(str),
        "Pen": _get(pen_col).astype(str),
        "Time1_raw": df[time1_col],
        "Time2_raw": _get(time2_col),
        "BW_1_kg": safe_numeric(_get(bw1_col)),
        "BW_2_kg": safe_numeric(_get(bw2_col)),
        "BW_3_kg": safe_numeric(_get(bw3_col)),
        "BW_kg": safe_numeric(_get(mean_col)),
        "Source_File": df.get("_Source_File", None),
        "Source_Sheet": df.get("_Source_Sheet", None),
    })

    # BW_kg 兜底：没有"平均体重"列时，用 3 个单次体重的均值
    if out["BW_kg"].isna().all():
        triple = out[["BW_1_kg", "BW_2_kg", "BW_3_kg"]]
        out["BW_kg"] = triple.mean(axis=1, skipna=True)

    out["Time1"] = parse_datetime_flexible(out["Time1_raw"])
    out["Time2"] = parse_datetime_flexible(out["Time2_raw"])
    out["Date"] = out["Time1"].dt.date
    return out


def make_experimental_week(dt_series, start_dt, week_length):
    dt_series = pd.to_datetime(dt_series)
    start_dt = pd.to_datetime(start_dt)
    days = (dt_series - start_dt).dt.total_seconds() / 86400
    day_num = np.floor(days).astype(int) + 1
    week_num = ((day_num - 1) // week_length) + 1
    return pd.DataFrame({
        "Experimental_Day": day_num,
        "Week": "Week " + week_num.astype(str),
        "Week_Number": week_num,
    })


def remove_weekly_outliers(df, value_col, week_col="Week_Number", sd_multiplier=3):
    stats_df = df.groupby(week_col)[value_col].agg(
        N=lambda s: s.notna().sum(),
        Mean=lambda s: s.mean(),
        SD=lambda s: s.std(ddof=1) if s.notna().sum() > 1 else np.nan,
    ).reset_index()
    stats_df["Lower"] = stats_df.apply(
        lambda r: -np.inf if pd.isna(r["SD"]) else (r["Mean"] if r["SD"] == 0 else r["Mean"] - sd_multiplier * r["SD"]),
        axis=1)
    stats_df["Upper"] = stats_df.apply(
        lambda r: np.inf if pd.isna(r["SD"]) else (r["Mean"] if r["SD"] == 0 else r["Mean"] + sd_multiplier * r["SD"]),
        axis=1)
    merged = df.merge(stats_df[[week_col, "Lower", "Upper"]], on=week_col, how="left")
    merged["QC_Outlier"] = ((merged[value_col] < merged["Lower"]) |
                            (merged[value_col] > merged["Upper"])) & merged[value_col].notna()
    return merged, stats_df


def drop_existing_columns(df, cols):
    """稳健版 drop：只删除 df 中实际存在的列，避免 KeyError。"""
    return df.drop(columns=[c for c in cols if c in df.columns])


def calculate_bouts(feed_clean):
    df = feed_clean.sort_values(["Animal_ID", "Time1", "Time2"]).copy()
    df["Feeding_Bout_Number"] = df.groupby("Animal_ID").cumcount() + 1
    df["Previous_Feeding_Time"] = df.groupby("Animal_ID")["Time1"].shift(1)
    df["IMI_sec"] = (df["Time1"] - df["Previous_Feeding_Time"]).dt.total_seconds()
    df["FR_g_sec"] = np.where(df["Feed_Duration_sec"] > 0,
                              df["Feed_Intake_g"] / df["Feed_Duration_sec"], np.nan)
    df["Hour"] = df["Time1"].dt.hour
    df["Minute"] = df["Time1"].dt.minute
    df["Time_of_Day_h"] = df["Hour"] + df["Minute"] / 60 + df["Time1"].dt.second / 3600
    return df.reset_index(drop=True)


def summarize_feeding(bout):
    def _safe_mean(s):
        s = s[np.isfinite(s)]
        return s.mean() if len(s) else np.nan

    g = bout.groupby("Animal_ID")
    out = pd.DataFrame({
        "TFB": g.size(),
        "FI_g": g["Feed_Intake_g"].sum(min_count=1),
        "AMS_g": g["Feed_Intake_g"].apply(_safe_mean),
        "TFD_sec": g["Feed_Duration_sec"].sum(min_count=1),
        "AFBD_sec": g["Feed_Duration_sec"].apply(_safe_mean),
        "IMI_sec": g["IMI_sec"].apply(_safe_mean),
        "FR_g_sec": g["FR_g_sec"].apply(_safe_mean),
        "Feed_Days": g["Date"].nunique(),
        "Feed_Start_Date": g["Date"].min(),
        "Feed_End_Date": g["Date"].max(),
    }).reset_index()
    return out


def calculate_production(bw_clean, feeding):
    bw = bw_clean.sort_values(["Animal_ID", "Time1"])
    bw_agg = bw.groupby("Animal_ID").agg(
        Group=("Group", "first"),
        House=("House", "first"),
        Pen=("Pen", "first"),
        IBW_kg=("BW_kg", "first"),
        FBW_kg=("BW_kg", "last"),
        Initial_Date=("Date", "first"),
        Final_Date=("Date", "last"),
        Initial_Time=("Time1", "first"),
        Final_Time=("Time1", "last"),
        BW_Records=("BW_kg", "size"),
    ).reset_index()
    bw_agg["Days"] = (pd.to_datetime(bw_agg["Final_Time"]) - pd.to_datetime(bw_agg["Initial_Time"])).dt.total_seconds() / 86400
    bw_agg["Gain_kg"] = bw_agg["FBW_kg"] - bw_agg["IBW_kg"]
    bw_agg["ADG_g"] = np.where(bw_agg["Days"] > 0, bw_agg["Gain_kg"] * 1000 / bw_agg["Days"], np.nan)
    bw_agg["MBW"] = ((bw_agg["IBW_kg"] + bw_agg["FBW_kg"]) / 2) ** 0.75

    prod = bw_agg.merge(feeding, on="Animal_ID", how="outer")
    prod["Group"] = prod["Group"].fillna("Unclassified").replace("", "Unclassified")
    prod["ADFI_g"] = np.where(prod["Feed_Days"] > 0, prod["FI_g"] / prod["Feed_Days"], np.nan)
    prod["FCR"] = np.where((prod["Gain_kg"] > 0) & prod["FI_g"].notna(),
                           prod["FI_g"] / (prod["Gain_kg"] * 1000), np.nan)
    prod["Feed_Efficiency"] = np.where((prod["FI_g"] > 0) & prod["Gain_kg"].notna(),
                                       prod["Gain_kg"] * 1000 / prod["FI_g"], np.nan)

    fit_dat = prod[["ADFI_g", "MBW", "ADG_g"]].dropna()
    prod["RFI"] = np.nan
    if len(fit_dat) >= 5 and fit_dat["MBW"].nunique() > 1 and fit_dat["ADG_g"].nunique() > 1:
        try:
            import statsmodels.api as sm
            X = sm.add_constant(fit_dat[["MBW", "ADG_g"]])
            model = sm.OLS(fit_dat["ADFI_g"], X).fit()
            pred = model.predict(X)
            resid = fit_dat["ADFI_g"] - pred
            idx_map = prod.dropna(subset=["ADFI_g", "MBW", "ADG_g"]).index
            prod.loc[idx_map, "RFI"] = resid.values
        except ImportError:
            pass
    prod = prod.loc[:, ~prod.columns.duplicated()].copy()
    return prod


def make_hff_lff(production, method, quantile):
    prod = production.copy()
    prod["Feed_Frequency_Group"] = None
    dat = prod[prod["TFB"].notna()]
    if len(dat) < 2:
        return prod
    if method == "median":
        cutoff = dat["TFB"].median()
        prod["Feed_Frequency_Group"] = np.where(prod["TFB"] >= cutoff, "HFF", "LFF")
    else:
        q_lo = dat["TFB"].quantile(quantile)
        q_hi = dat["TFB"].quantile(1 - quantile)
        prod["Feed_Frequency_Group"] = np.select(
            [prod["TFB"] >= q_hi, prod["TFB"] <= q_lo],
            ["HFF", "LFF"], default="Middle")
    return prod


def hff_test(prod, variable):
    x = prod[prod["Feed_Frequency_Group"].isin(["HFF", "LFF"])].dropna(subset=[variable])
    if len(x) < 4 or x["Feed_Frequency_Group"].nunique() < 2:
        return {"指标": variable, "HFF_N": np.nan, "LFF_N": np.nan,
                "HFF_Median": np.nan, "LFF_Median": np.nan,
                "Mann_Whitney_P": np.nan, "T_test_P": np.nan}
    hff = x.loc[x["Feed_Frequency_Group"] == "HFF", variable]
    lff = x.loc[x["Feed_Frequency_Group"] == "LFF", variable]
    try:
        mw_p = stats.mannwhitneyu(hff, lff, alternative="two-sided").pvalue
    except Exception:
        mw_p = np.nan
    try:
        t_p = stats.ttest_ind(hff, lff, equal_var=False, nan_policy="omit").pvalue
    except Exception:
        t_p = np.nan
    return {"指标": variable, "HFF_N": len(hff), "LFF_N": len(lff),
            "HFF_Median": np.nanmedian(hff), "LFF_Median": np.nanmedian(lff),
            "Mann_Whitney_P": mw_p, "T_test_P": t_p}


def calculate_weekly_fcr(feed_clean, bw_clean):
    weekly_feed = feed_clean.groupby(["Animal_ID", "Week_Number"])["Feed_Intake_g"].sum(min_count=1).reset_index()
    weekly_feed.columns = ["Animal_ID", "Week_Number", "Weekly_FI_g"]

    bw = bw_clean.sort_values(["Animal_ID", "Time1"])
    weekly_bw = bw.groupby(["Animal_ID", "Week_Number"]).agg(
        Week_Initial_BW=("BW_kg", "first"),
        Week_Final_BW=("BW_kg", "last"),
    ).reset_index()
    weekly_bw["Weekly_Gain_kg"] = weekly_bw["Week_Final_BW"] - weekly_bw["Week_Initial_BW"]

    out = weekly_feed.merge(weekly_bw, on=["Animal_ID", "Week_Number"], how="outer")
    out["Week"] = "Week " + out["Week_Number"].astype(str)
    out["Weekly_FCR"] = np.where(out["Weekly_Gain_kg"] > 0,
                                 out["Weekly_FI_g"] / (out["Weekly_Gain_kg"] * 1000), np.nan)
    return out


def calculate_correlations(dat):
    """强制单列 + 去重列名，防止 'truth value of a Series is ambiguous'。"""
    vars_ = ["IBW_kg", "FBW_kg", "Gain_kg", "TFB", "FI_g", "AMS_g",
             "TFD_sec", "AFBD_sec", "IMI_sec", "FR_g_sec", "FCR", "RFI"]

    dat = dat.loc[:, ~dat.columns.duplicated()].copy()
    vars_ = [v for v in vars_ if v in dat.columns]

    n = len(vars_)
    r_mat = np.full((n, n), np.nan)
    p_mat = np.full((n, n), np.nan)

    for i in range(n):
        for j in range(n):
            try:
                si = dat[vars_[i]]
                sj = dat[vars_[j]]
                if isinstance(si, pd.DataFrame):
                    si = si.iloc[:, 0]
                if isinstance(sj, pd.DataFrame):
                    sj = sj.iloc[:, 0]
                tmp = pd.concat([si, sj], axis=1).dropna()
                tmp.columns = ["_a", "_b"]
                if (len(tmp) >= 3
                        and tmp["_a"].nunique() > 1
                        and tmp["_b"].nunique() > 1):
                    r, p = stats.spearmanr(tmp["_a"], tmp["_b"])
                    r_mat[i, j] = float(r)
                    p_mat[i, j] = float(p)
            except Exception:
                pass
    return pd.DataFrame(r_mat, index=vars_, columns=vars_), pd.DataFrame(p_mat, index=vars_, columns=vars_)


def safe_mean(s):
    s = s[np.isfinite(s)]
    return s.mean() if len(s) else np.nan


def safe_median(s):
    s = s[np.isfinite(s)]
    return np.median(s) if len(s) else np.nan


def safe_sd(s):
    s = s[np.isfinite(s)]
    return s.std(ddof=1) if len(s) >= 2 else np.nan


# ============================================================
# 创新行为指标
# ============================================================
def compute_daynight(bout, day_start_min, day_end_min):
    def is_day(t):
        mins = t.hour * 60 + t.minute
        if day_start_min < day_end_min:
            return day_start_min <= mins < day_end_min
        elif day_start_min > day_end_min:
            return mins >= day_start_min or mins < day_end_min
        return False

    bout = bout.copy()
    bout["_Is_Day"] = bout["Time1"].apply(is_day)
    rows = []
    for aid, g in bout.groupby("Animal_ID"):
        day = g[g["_Is_Day"]]
        night = g[~g["_Is_Day"]]
        day_fi = day["Feed_Intake_g"].sum(min_count=1)
        night_fi = night["Feed_Intake_g"].sum(min_count=1)
        total_fi = (day_fi or 0) + (night_fi or 0)
        rows.append({
            "Animal_ID": aid,
            "Day_FI_g": day_fi if not pd.isna(day_fi) else 0,
            "Night_FI_g": night_fi if not pd.isna(night_fi) else 0,
            "Day_Bouts": len(day),
            "Night_Bouts": len(night),
            "Day_Avg_Meal": safe_mean(day["Feed_Intake_g"].values),
            "Night_Avg_Meal": safe_mean(night["Feed_Intake_g"].values),
            "Day_Avg_Duration": safe_mean(day["Feed_Duration_sec"].values),
            "Night_Avg_Duration": safe_mean(night["Feed_Duration_sec"].values),
            "Total_FI_g": total_fi,
            "Total_Bouts": len(g),
        })
    df = pd.DataFrame(rows)
    df["Day_FI_Ratio"] = np.where(df["Total_FI_g"] > 0, df["Day_FI_g"] / df["Total_FI_g"], np.nan)
    df["Day_Bout_Ratio"] = np.where(df["Total_Bouts"] > 0, df["Day_Bouts"] / df["Total_Bouts"], np.nan)
    df["Day_Night_FI_Ratio"] = np.where(df["Night_FI_g"] > 0, df["Day_FI_g"] / df["Night_FI_g"], np.nan)
    df["Day_Night_Bout_Ratio"] = np.where(df["Night_Bouts"] > 0, df["Day_Bouts"] / df["Night_Bouts"], np.nan)
    return df


def compute_behavior_cv(bout, min_bouts_single=20, min_days_daily=3):
    rows = []
    for aid, g in bout.groupby("Animal_ID"):
        n_bouts = len(g)
        n_days = g["Date"].nunique()
        cv_dur = safe_sd(g["Feed_Duration_sec"].values) / safe_mean(g["Feed_Duration_sec"].values) \
            if n_bouts >= min_bouts_single and safe_mean(g["Feed_Duration_sec"].values) > 0 else np.nan
        cv_fr = safe_sd(g["FR_g_sec"].values) / safe_mean(g["FR_g_sec"].values) \
            if n_bouts >= min_bouts_single and safe_mean(g["FR_g_sec"].values) > 0 else np.nan
        imi = g["IMI_sec"].dropna().values
        if n_bouts >= min_bouts_single and len(imi) and safe_median(imi) > 0:
            q75, q25 = np.percentile(imi, [75, 25])
            rcv = (q75 - q25) / safe_median(imi)
        else:
            rcv = np.nan

        daily = g.groupby("Date").agg(Daily_Bouts=("Feed_Intake_g", "size"),
                                       Daily_FI=("Feed_Intake_g", "sum"),
                                       Daily_TFD=("Feed_Duration_sec", "sum"))
        n_days_daily = len(daily)
        cv_db = safe_sd(daily["Daily_Bouts"].values) / safe_mean(daily["Daily_Bouts"].values) \
            if n_days_daily >= min_days_daily and safe_mean(daily["Daily_Bouts"].values) > 0 else np.nan
        cv_df = safe_sd(daily["Daily_FI"].values) / safe_mean(daily["Daily_FI"].values) \
            if n_days_daily >= min_days_daily and safe_mean(daily["Daily_FI"].values) > 0 else np.nan
        cv_dt = safe_sd(daily["Daily_TFD"].values) / safe_mean(daily["Daily_TFD"].values) \
            if n_days_daily >= min_days_daily and safe_mean(daily["Daily_TFD"].values) > 0 else np.nan

        rows.append({
            "Animal_ID": aid, "N_Bouts": n_bouts, "N_Days": n_days,
            "CV_Duration": cv_dur, "CV_FR": cv_fr, "Robust_CV_IMI": rcv,
            "N_Days_Daily": n_days_daily,
            "CV_Daily_Bouts": cv_db, "CV_Daily_FI": cv_df, "CV_Daily_TFD": cv_dt,
        })
    return pd.DataFrame(rows)


def compute_cosinor(bout, min_bouts=20):
    rows = []
    for aid, g in bout.groupby("Animal_ID"):
        hourly = g.groupby("Hour").size().reindex(range(24), fill_value=0)
        y = hourly.values.astype(float)
        if y.sum() < min_bouts:
            rows.append({"Animal_ID": aid, "Cosinor_M": np.nan, "Cosinor_A": np.nan,
                         "Peak_Hour": np.nan, "Cosinor_R2": np.nan})
            continue
        t = np.arange(24)
        X = np.column_stack([np.ones(24), np.cos(2 * np.pi * t / 24), np.sin(2 * np.pi * t / 24)])
        try:
            coef, *_ = np.linalg.lstsq(X, y, rcond=None)
            pred = X @ coef
            ss_res = np.sum((y - pred) ** 2)
            ss_tot = np.sum((y - y.mean()) ** 2)
            r2 = 1 - ss_res / ss_tot if ss_tot > 0 else np.nan
            b0, bc, bs = coef
            A = math.hypot(bc, bs)
            peak = (math.atan2(bs, bc) * 24 / (2 * math.pi)) % 24
            rows.append({"Animal_ID": aid, "Cosinor_M": b0, "Cosinor_A": A,
                         "Peak_Hour": peak, "Cosinor_R2": r2})
        except Exception:
            rows.append({"Animal_ID": aid, "Cosinor_M": np.nan, "Cosinor_A": np.nan,
                         "Peak_Hour": np.nan, "Cosinor_R2": np.nan})
    return pd.DataFrame(rows)


def compute_fano(bout):
    rows = []
    for aid, g in bout.groupby("Animal_ID"):
        cells = g.groupby(["Date", "Hour"]).size()
        dates = g["Date"].unique()
        full_index = pd.MultiIndex.from_product([dates, range(24)], names=["Date", "Hour"])
        cells = cells.reindex(full_index, fill_value=0)
        n = cells.values
        mean_n = n.mean()
        var_n = n.var(ddof=1) if len(n) > 1 else np.nan
        fano = var_n / mean_n if (len(n) > 1 and mean_n > 0) else np.nan
        rows.append({"Animal_ID": aid, "N_Hour_Cells": len(n),
                     "Mean_Bouts_Per_Hour": mean_n,
                     "Var_Bouts_Per_Hour": var_n, "Fano": fano})
    return pd.DataFrame(rows)


def compute_innovation_corr(dat, innovation_vars):
    if dat is None or len(dat) == 0:
        return pd.DataFrame()
    dat = dat.loc[:, ~dat.columns.duplicated()].copy()
    perf_vars = [v for v in ["ADG_g", "FCR", "RFI", "ADFI_g", "Feed_Efficiency"] if v in dat.columns]
    inn_vars = [v for v in innovation_vars if v in dat.columns]
    if not perf_vars or not inn_vars:
        return pd.DataFrame()
    rows = []
    for iv in inn_vars:
        for pv in perf_vars:
            try:
                si = dat[iv]
                sj = dat[pv]
                if isinstance(si, pd.DataFrame):
                    si = si.iloc[:, 0]
                if isinstance(sj, pd.DataFrame):
                    sj = sj.iloc[:, 0]
                tmp = pd.concat([si, sj], axis=1).dropna()
                tmp.columns = ["_a", "_b"]
                if len(tmp) >= 5 and tmp["_a"].nunique() > 1 and tmp["_b"].nunique() > 1:
                    r, p = stats.spearmanr(tmp["_a"], tmp["_b"])
                    rows.append({"Innovation": iv, "Performance": pv,
                                 "Spearman_r": r, "P_value": p, "N": len(tmp)})
                else:
                    rows.append({"Innovation": iv, "Performance": pv,
                                 "Spearman_r": np.nan, "P_value": np.nan, "N": len(tmp)})
            except Exception:
                rows.append({"Innovation": iv, "Performance": pv,
                             "Spearman_r": np.nan, "P_value": np.nan, "N": 0})
    return pd.DataFrame(rows)


# ============================================================
# 绘图工具
# ============================================================
def fig_to_png_bytes(fig, width=1200, height=800, scale=2):
    try:
        return fig.to_image(format="png", width=width, height=height, scale=scale)
    except Exception as e:
        raise RuntimeError(
            "PNG 导出失败：Plotly 导出图片依赖 kaleido。\n"
            "请在终端运行：pip install -U kaleido -i https://pypi.tuna.tsinghua.edu.cn/simple\n"
            f"原始错误：{e}"
        )


def dl_button(fig, key, filename):
    try:
        png = fig_to_png_bytes(fig)
        st.download_button("下载PNG", data=png, file_name=filename,
                           mime="image/png", key=key)
    except Exception as e:
        st.button("下载PNG（需先安装 kaleido）", key=key, disabled=True, help=str(e))


# ============================================================
# Excel 导出工具函数（供侧边栏和导出页共用）
# ============================================================
EXPORT_MAP_KEYS = {
    "qc": "QC 结果",
    "clean": "清洗后采食/体重数据",
    "bout": "完整单次采食数据",
    "individual": "个体综合指标",
    "hff": "HFF / LFF 及检验",
    "rhythm": "采食节律",
    "weekly_fcr": "每周 FCR",
    "correlation": "Spearman 相关",
    "innovation": "创新行为指标",
}


def _collect_export_sheets(result_obj, key):
    """返回 [(sheet_name, df), ...]，按导出类别分组"""
    mapping = {
        "qc": [
            ("QC_Summary", result_obj.get("qc_summary")),
            ("DateTime_QC", result_obj.get("datetime_qc")),
            ("Read_QC", result_obj.get("read_qc")),
            ("Week_Summary", result_obj.get("week_summary")),
            ("NA_QC", result_obj.get("na_qc")),
            ("Feed_Weekly_QC", result_obj.get("feed_week_stats")),
            ("BW_Weekly_QC", result_obj.get("bw_week_stats")),
            ("Feed_Outliers", result_obj.get("feed_outlier")),
            ("BW_Outliers", result_obj.get("bw_outlier")),
        ],
        "clean": [
            ("Feed_Clean", result_obj.get("feed_clean")),
            ("BW_Clean", result_obj.get("bw_clean")),
        ],
        "bout": [("Feeding_Bouts", result_obj.get("bout"))],
        "individual": [("Individual", result_obj.get("individual"))],
        "hff": [
            ("HFF_LFF", result_obj.get("hff")),
            ("HFF_LFF_Group_Summary", result_obj.get("hff_group_summary")),
            ("HFF_LFF_Test", result_obj.get("hff_tests")),
        ],
        "rhythm": [
            ("Daily_Bouts", result_obj.get("daily_bouts")),
            ("Weekly_Bouts", result_obj.get("weekly_bouts")),
            ("Feeding_Rhythm", result_obj.get("rhythm")),
        ],
        "weekly_fcr": [
            ("Weekly_FCR", result_obj.get("weekly_fcr")),
            ("Weekly_FCR_Test", result_obj.get("weekly_fcr_tests")),
        ],
        "correlation": [
            ("Spearman_R", result_obj.get("correlation")),
            ("Spearman_P", result_obj.get("correlation_p")),
        ],
        "innovation": [
            ("DayNight_Distribution", result_obj.get("daynight_table")),
            ("Behavior_CV", result_obj.get("cv_table")),
            ("Cosinor_Fit", result_obj.get("cosinor_table")),
            ("Fano_Index", result_obj.get("fano_table")),
            ("Innovation_Merged", result_obj.get("innovation_merged")),
        ],
    }
    return mapping.get(key, [])


def _df_to_ws(wb, sheet_name, df):
    if df is None or len(df) == 0:
        return False
    ws = wb.create_sheet(str(sheet_name)[:31])
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    return True


def build_excel(result_obj, keys, params_info=None):
    """按 keys 列表构建 Excel 文件，返回 bytes。params_info 为 dict 时，会写入 Parameters 页。"""
    wb = Workbook()
    wb.remove(wb.active)

    for key in keys:
        for name, df in _collect_export_sheets(result_obj, key):
            if key == "correlation" and df is not None and len(df) > 0:
                _df_to_ws(wb, name, df.reset_index().rename(columns={"index": "Indicator"}))
            else:
                _df_to_ws(wb, name, df)

    if params_info:
        ws = wb.create_sheet("Parameters")
        rows = [["Parameter", "Value"]] + [[k, v] for k, v in params_info.items()]
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ============================================================
# 侧边栏
# ============================================================
st.sidebar.title("肉鸭综合分析工具 V5.0")

st.sidebar.header("① 数据上传")
feed_files = st.sidebar.file_uploader("采食原始数据（可上传 N 个 Excel）",
                                       type=["xlsx", "xls"], accept_multiple_files=True)
weight_files = st.sidebar.file_uploader("体重原始数据（可上传 N 个 Excel）",
                                         type=["xlsx", "xls"], accept_multiple_files=True)

st.sidebar.header("② 实验时间参数")
training_date = st.sidebar.date_input("训练结束日期", value=datetime.today().date())
training_time = st.sidebar.text_input("训练结束时间（HH:MM:SS）", value="00:00:00")
remove_before = st.sidebar.checkbox("删除训练结束时间以前的数据", value=True)
week_length = st.sidebar.number_input("每周天数", min_value=1, max_value=60, value=7, step=1)
rhythm_hours = st.sidebar.number_input("采食节律时间间隔（小时）", min_value=1, max_value=6, value=1, step=1)

st.sidebar.header("②b 昼夜时段")
day_start = st.sidebar.text_input("白天开始（HH:MM）", value="06:00")
day_end = st.sidebar.text_input("白天结束（HH:MM）", value="18:00")

st.sidebar.header("③ 异常值参数")
feed_sd = st.sidebar.number_input("采食时长异常值：μ ± kσ", min_value=0.5, max_value=10.0, value=3.0, step=0.5)
bw_sd = st.sidebar.number_input("平均体重异常值：μ ± kσ", min_value=0.5, max_value=10.0, value=3.0, step=0.5)

st.sidebar.header("④ HFF / LFF 分组")
hff_method = st.sidebar.selectbox("分组依据", ["median", "iqr"],
                                   format_func=lambda x: "TFB 中位数" if x == "median" else "TFB 上下四分位数")
hff_quantile = st.sidebar.number_input("TFB 分位数阈值（IQR模式）", min_value=0.05, max_value=0.45,
                                        value=0.25, step=0.05)

st.sidebar.header("⑤ 分析模块")
modules_all = {
    "feeding": "采食行为（TFB/FI/AMS/TFD/AFBD/IMI/FR）",
    "production": "生产性能（IBW/FBW/Gain/ADG/ADFI/MBW/FCR/RFI）",
    "hff": "HFF / LFF 分组与差异检验",
    "rhythm": "采食节律（每日/每周/24 h）",
    "weekly_fcr": "每周 FCR",
    "correlation": "12指标 Spearman 相关",
    "bout_kw": "单次采食周间差异（Kruskal-Wallis）",
    "innovation": "创新行为指标（昼夜/CV/余弦/Fano）",
}
selected_modules = st.sidebar.multiselect(
    "选择模块", options=list(modules_all.keys()),
    default=list(modules_all.keys()),
    format_func=lambda k: modules_all[k]
)

st.sidebar.header("⑥ Excel 导出项目")
export_all = {
    "qc": "QC 结果",
    "clean": "清洗后采食/体重数据",
    "bout": "完整单次采食数据",
    "individual": "个体综合指标",
    "hff": "HFF / LFF 及检验",
    "rhythm": "采食节律",
    "weekly_fcr": "每周 FCR",
    "correlation": "Spearman 相关",
    "innovation": "创新行为指标",
}
selected_exports = st.sidebar.multiselect(
    "选择导出项", options=list(export_all.keys()),
    default=["qc", "individual", "hff", "rhythm", "weekly_fcr", "correlation", "innovation"],
    format_func=lambda k: export_all[k]
)

run_btn = st.sidebar.button("开始分析", type="primary", use_container_width=True)


# ============================================================
# Session state
# ============================================================
if "result" not in st.session_state:
    st.session_state.result = {}


def _params_info():
    return {
        "Export_Time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Training_End": f"{training_date} {training_time}",
        "Remove_Before": str(remove_before),
        "Week_Length_Days": str(week_length),
        "Feed_SD_Multiplier": str(feed_sd),
        "BW_SD_Multiplier": str(bw_sd),
        "HFF_Method": hff_method,
        "HFF_Quantile": str(hff_quantile),
        "Day_Start": parse_clock_hhmm(day_start),
        "Day_End": parse_clock_hhmm(day_end),
        "Analysis_Modules": ",".join(selected_modules),
        "Export_Sections": ",".join(selected_exports),
    }


# ---------- 侧边栏：一键导出（分析完成后才显示） ----------
if st.session_state.result:
    st.sidebar.header("⑦ 一键导出")
    _ts = datetime.today().strftime("%Y%m%d_%H%M%S")
    try:
        _keys = list(selected_exports) if selected_exports else list(export_all.keys())
        _xls_bytes = build_excel(st.session_state.result, _keys, params_info=_params_info())
        st.sidebar.download_button(
            label=f"📥 导出所选 Excel（{len(_keys)} 个类别）",
            data=_xls_bytes,
            file_name=f"肉鸭分析_V5.0_{_ts}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key="sidebar_dl_all",
        )
        st.sidebar.caption(f"文件大小约 {len(_xls_bytes)/1024:.1f} KB")
    except Exception as _e:
        st.sidebar.error(f"侧边栏导出失败：{_e}")


def run_analysis():
    res = {}
    feed_bytes = [(f.name, f.read()) for f in feed_files]
    weight_bytes = [(f.name, f.read()) for f in weight_files]

    feed_original, qc_feed = read_matching_excel_files(feed_bytes, "feed")
    weight_original, qc_weight = read_matching_excel_files(weight_bytes, "weight")

    qc_parts = [d for d in (qc_feed, qc_weight) if d is not None and len(d) > 0]
    res["read_qc"] = pd.concat(qc_parts, ignore_index=True) if qc_parts else pd.DataFrame()

    if len(feed_original) == 0 or len(weight_original) == 0:
        st.session_state.result = res
        msgs = []
        if len(feed_original) == 0:
            msgs.append("❌ 采食数据没有匹配到工作表")
        if len(weight_original) == 0:
            msgs.append("❌ 体重数据没有匹配到工作表")
        msgs.append("请切到上方『QC』标签页，查看『读取工作表记录』表格里的『状态』和『说明』列。")
        raise ValueError(" | ".join(msgs))

    feed_raw = standardize_feed(feed_original)
    bw_raw = standardize_bw(weight_original)

    if len(feed_raw) > 0 and feed_raw["Feed_Duration_sec"].isna().all():
        st.warning("⚠️ 采食时长字段未识别，FR/TFD/AFBD 等指标将为空。")
    if len(feed_raw) > 0 and feed_raw["Feed_Intake_g"].isna().all():
        st.warning("⚠️ 处理后采食量字段未识别，FI/ADFI/FCR 等指标将为空。")

    res["datetime_qc"] = pd.DataFrame({
        "数据集": ["采食数据", "体重数据"],
        "Time1_解析成功率(%)": [
            round(feed_raw["Time1"].notna().mean() * 100, 2) if len(feed_raw) else np.nan,
            round(bw_raw["Time1"].notna().mean() * 100, 2) if len(bw_raw) else np.nan,
        ],
    })

    res["na_qc"] = pd.DataFrame({
        "数据集": ["采食数据"] * 5 + ["体重数据"] * 7,
        "字段": ["耳标", "记录时间1", "记录时间2", "处理后采食量(g)", "采食时长",
               "耳标", "记录时间1", "记录时间2", "体重1(kg)", "体重2(kg)", "体重3(kg)", "平均体重(kg)"],
        "NA记录数": [
            feed_raw["Animal_ID"].isna().sum() + (feed_raw["Animal_ID"].astype(str).str.strip() == "").sum(),
            feed_raw["Time1"].isna().sum(), feed_raw["Time2"].isna().sum(),
            feed_raw["Feed_Intake_g"].isna().sum(), feed_raw["Feed_Duration_sec"].isna().sum(),
            bw_raw["Animal_ID"].isna().sum() + (bw_raw["Animal_ID"].astype(str).str.strip() == "").sum(),
            bw_raw["Time1"].isna().sum(), bw_raw["Time2"].isna().sum(),
            bw_raw["BW_1_kg"].isna().sum(), bw_raw["BW_2_kg"].isna().sum(),
            bw_raw["BW_3_kg"].isna().sum(), bw_raw["BW_kg"].isna().sum(),
        ],
    })

    feed_complete = feed_raw.dropna(subset=["Time1"]).copy()
    feed_complete = feed_complete[feed_complete["Animal_ID"].astype(str).str.strip() != ""]
    bw_complete = bw_raw.dropna(subset=["Time1", "BW_kg"]).copy()
    bw_complete = bw_complete[bw_complete["Animal_ID"].astype(str).str.strip() != ""]

    feed_deleted_na = len(feed_raw) - len(feed_complete)
    bw_deleted_na = len(bw_raw) - len(bw_complete)

    try:
        hh, mm, ss = map(int, training_time.split(":"))
    except Exception:
        hh, mm, ss = 0, 0, 0
    end_dt = pd.Timestamp(datetime.combine(training_date, datetime.min.time()).replace(hour=hh, minute=mm, second=ss))

    feed_before_time = len(feed_complete)
    bw_before_time = len(bw_complete)
    if remove_before:
        feed_time = feed_complete[feed_complete["Time1"] >= end_dt].copy()
        bw_time = bw_complete[bw_complete["Time1"] >= end_dt].copy()
    else:
        feed_time, bw_time = feed_complete.copy(), bw_complete.copy()
    feed_deleted_time = feed_before_time - len(feed_time)
    bw_deleted_time = bw_before_time - len(bw_time)

    feed_time = pd.concat([feed_time.reset_index(drop=True),
                           make_experimental_week(feed_time["Time1"], end_dt, week_length).reset_index(drop=True)],
                          axis=1)
    bw_time = pd.concat([bw_time.reset_index(drop=True),
                         make_experimental_week(bw_time["Time1"], end_dt, week_length).reset_index(drop=True)],
                        axis=1)

    feed_marked, feed_week_stats = remove_weekly_outliers(feed_time, "Feed_Duration_h", "Week_Number", feed_sd)
    feed_outlier = feed_marked[feed_marked["QC_Outlier"]].copy()
    feed_clean = drop_existing_columns(
        feed_marked[~feed_marked["QC_Outlier"]],
        ["N", "Mean", "SD", "Lower", "Upper", "QC_Outlier"]
    ).reset_index(drop=True)

    bw_marked, bw_week_stats = remove_weekly_outliers(bw_time, "BW_kg", "Week_Number", bw_sd)
    bw_outlier = bw_marked[bw_marked["QC_Outlier"]].copy()
    bw_clean = drop_existing_columns(
        bw_marked[~bw_marked["QC_Outlier"]],
        ["N", "Mean", "SD", "Lower", "Upper", "QC_Outlier"]
    ).reset_index(drop=True)

    res["feed_clean"] = feed_clean
    res["bw_clean"] = bw_clean
    res["feed_outlier"] = feed_outlier
    res["bw_outlier"] = bw_outlier
    res["feed_week_stats"] = feed_week_stats
    res["bw_week_stats"] = bw_week_stats

    all_weeks = sorted(set(feed_time["Week_Number"].dropna().astype(int)) |
                        set(feed_clean["Week_Number"].dropna().astype(int)) |
                        set(bw_time["Week_Number"].dropna().astype(int)) |
                        set(bw_clean["Week_Number"].dropna().astype(int)))
    if all_weeks:
        max_w = max(all_weeks)
        week_df = pd.DataFrame({"Week_Number": range(1, max_w + 1)})
        f_before = feed_time.groupby("Week_Number").size().rename("Feed_Before_QC")
        f_after = feed_clean.groupby("Week_Number").size().rename("Feed_After_QC")
        b_before = bw_time.groupby("Week_Number").size().rename("BW_Before_QC")
        b_after = bw_clean.groupby("Week_Number").size().rename("BW_After_QC")
        week_df = week_df.merge(f_before, on="Week_Number", how="left") \
                         .merge(f_after, on="Week_Number", how="left") \
                         .merge(b_before, on="Week_Number", how="left") \
                         .merge(b_after, on="Week_Number", how="left").fillna(0)
        week_df["Week"] = "Week " + week_df["Week_Number"].astype(str)
        week_df["Feed_QC_Removed"] = week_df["Feed_Before_QC"] - week_df["Feed_After_QC"]
        week_df["BW_QC_Removed"] = week_df["BW_Before_QC"] - week_df["BW_After_QC"]
        res["week_summary"] = week_df
    else:
        res["week_summary"] = pd.DataFrame()

    need_bout = any(m in selected_modules for m in ["feeding", "production", "hff", "rhythm", "bout_kw", "innovation"])
    need_production = any(m in selected_modules for m in ["production", "hff", "correlation", "innovation"])

    bout = feeding = production = None
    if need_bout:
        bout = calculate_bouts(feed_clean)
        feeding = summarize_feeding(bout)
        res["bout"] = bout
        res["feeding"] = feeding

    if need_production and feeding is not None:
        production = calculate_production(bw_clean, feeding)
        res["individual"] = production
        res["production"] = production
    elif need_production:
        empty_feeding = pd.DataFrame(columns=["Animal_ID", "TFB", "FI_g", "AMS_g", "TFD_sec",
                                               "AFBD_sec", "IMI_sec", "FR_g_sec", "Feed_Days"])
        production = calculate_production(bw_clean, empty_feeding)
        res["individual"] = production
        res["production"] = production

    if "hff" in selected_modules and production is not None and "TFB" in production.columns:
        prod_hff = make_hff_lff(production, hff_method, hff_quantile)
        res["individual"] = prod_hff
        res["production"] = prod_hff
        test_vars = ["TFB", "FI_g", "AMS_g", "TFD_sec", "AFBD_sec", "IMI_sec", "FR_g_sec",
                     "IBW_kg", "FBW_kg", "Gain_kg", "ADG_g", "ADFI_g", "FCR", "RFI"]
        test_vars = [v for v in test_vars if v in prod_hff.columns]
        res["hff_tests"] = pd.DataFrame([hff_test(prod_hff, v) for v in test_vars])

        grp_sum = prod_hff[prod_hff["Feed_Frequency_Group"].isin(["HFF", "LFF"])].groupby("Feed_Frequency_Group").agg(
            N=("Animal_ID", "size"),
            TFB_Median=("TFB", "median"),
            TFB_Mean=("TFB", "mean"),
            TFB_Min=("TFB", "min"),
            TFB_Max=("TFB", "max"),
        ).reset_index()
        res["hff_group_summary"] = grp_sum

        cols = ["Animal_ID", "Group", "House", "Pen", "Feed_Frequency_Group",
                "TFB", "FI_g", "AMS_g", "TFD_sec", "AFBD_sec", "IMI_sec", "FR_g_sec",
                "IBW_kg", "FBW_kg", "Gain_kg", "ADG_g", "ADFI_g", "MBW", "FCR", "RFI"]
        cols = [c for c in cols if c in prod_hff.columns]
        res["hff"] = prod_hff[cols].copy()

    if "rhythm" in selected_modules and bout is not None:
        daily = bout.dropna(subset=["Date"]).groupby(["Animal_ID", "Date"]).size().reset_index(name="Bouts_Per_Duck")
        daily_sum = daily.groupby("Date").agg(
            Total_Bouts=("Bouts_Per_Duck", "sum"),
            Animal_N=("Bouts_Per_Duck", "size"),
            Mean_Bouts_Per_Duck=("Bouts_Per_Duck", "mean"),
        ).reset_index().sort_values("Date")
        res["daily_bouts"] = daily_sum

        weekly = bout.dropna(subset=["Date"]).groupby(["Animal_ID", "Week_Number", "Week", "Date"]).size().reset_index(name="Bouts_Per_Duck")
        weekly_sum = weekly.groupby(["Week_Number", "Week", "Date"]).agg(
            Total_Bouts=("Bouts_Per_Duck", "sum"),
            Animal_N=("Bouts_Per_Duck", "size"),
            Mean_Bouts_Per_Duck=("Bouts_Per_Duck", "mean"),
        ).reset_index().sort_values(["Week_Number", "Date"])
        res["weekly_bouts"] = weekly_sum

        step = int(rhythm_hours)
        sub = bout[(bout["Week_Number"] >= 1) & (bout["Week_Number"] <= 5)].copy()
        sub["Hour_Block"] = (sub["Hour"] // step) * step
        rhythm = sub.groupby(["Week_Number", "Hour_Block"]).agg(
            Total_Bouts=("Animal_ID", "size"),
            Animal_N=("Animal_ID", "nunique"),
        ).reset_index()
        rhythm["Mean_Bouts_Per_Duck"] = rhythm["Total_Bouts"] / rhythm["Animal_N"].replace(0, np.nan)
        full_idx = pd.MultiIndex.from_product([range(1, 6), range(0, 24, step)],
                                               names=["Week_Number", "Hour_Block"])
        rhythm = rhythm.set_index(["Week_Number", "Hour_Block"]).reindex(full_idx, fill_value=0).reset_index()
        res["rhythm"] = rhythm

    if "innovation" in selected_modules and bout is not None:
        ds_min = clock_to_minutes(day_start)
        de_min = clock_to_minutes(day_end)
        daynight = compute_daynight(bout, ds_min, de_min)
        cv = compute_behavior_cv(bout)
        cosinor = compute_cosinor(bout)
        fano = compute_fano(bout)
        res["daynight_table"] = daynight
        res["cv_table"] = cv
        res["cosinor_table"] = cosinor
        res["fano_table"] = fano

        if production is not None:
            perf_cols = [c for c in ["Animal_ID", "ADG_g", "FCR", "RFI", "ADFI_g",
                                      "Feed_Efficiency", "TFB", "FI_g", "Gain_kg", "FBW_kg"]
                         if c in production.columns]
            merged = production[perf_cols] \
                .merge(daynight, on="Animal_ID", how="left") \
                .merge(cv, on="Animal_ID", how="left") \
                .merge(cosinor, on="Animal_ID", how="left") \
                .merge(fano, on="Animal_ID", how="left")
            merged = merged.loc[:, ~merged.columns.duplicated()].copy()
            res["innovation_merged"] = merged

    if "weekly_fcr" in selected_modules:
        wfcr = calculate_weekly_fcr(feed_clean, bw_clean)
        res["weekly_fcr"] = wfcr
        wfcr_valid = wfcr.dropna(subset=["Weekly_FCR"])
        if len(wfcr_valid) > 0 and wfcr_valid["Week_Number"].nunique() >= 2:
            groups = [g["Weekly_FCR"].values for _, g in wfcr_valid.groupby("Week_Number")]
            try:
                p = stats.kruskal(*groups).pvalue
            except Exception:
                p = np.nan
        else:
            p = np.nan
        res["weekly_fcr_tests"] = pd.DataFrame({"检验": ["Kruskal-Wallis"],
                                                 "指标": ["Weekly FCR"],
                                                 "P值": [p]})

    if "correlation" in selected_modules and production is not None:
        r, p = calculate_correlations(production)
        res["correlation"] = r
        res["correlation_p"] = p

    res["qc_summary"] = pd.DataFrame({
        "数据集": ["采食数据", "体重数据"],
        "原始记录数": [len(feed_raw), len(bw_raw)],
        "删除重要字段NA": [feed_deleted_na, bw_deleted_na],
        "删除训练结束前": [feed_deleted_time, bw_deleted_time],
        "删除周内异常值": [len(feed_outlier), len(bw_outlier)],
        "最终保留记录": [len(feed_clean), len(bw_clean)],
        "保留率(%)": [
            round(len(feed_clean) / len(feed_raw) * 100, 2) if len(feed_raw) else np.nan,
            round(len(bw_clean) / len(bw_raw) * 100, 2) if len(bw_raw) else np.nan,
        ],
    })
    return res


if run_btn:
    if not feed_files or not weight_files:
        st.error("请先上传采食和体重 Excel 文件。")
    else:
        with st.spinner("正在分析，请稍候……"):
            try:
                # 清空旧结果，避免异常残留
                st.session_state.result = {}
                st.session_state.result = run_analysis()
                st.success("分析完成！")
                st.info("💡 导出按钮已出现在左侧边栏『⑦ 一键导出』，也可以切到『结果导出』Tab 里下载分类文件。")
            except Exception as e:
                st.exception(e)


# ============================================================
# 主界面
# ============================================================
result = st.session_state.result

# 12 个 Tab（减少被折叠的可能）：把"结果导出"合并进"清洗后数据"后面的独立 Tab 保留，但总数控制在 12
tab_names = ["使用说明", "QC", "个体指标", "采食行为", "生产性能",
             "HFF / LFF", "采食节律", "创新行为指标", "单次采食 / 生长曲线",
             "每周 FCR / 相关性", "清洗后数据", "结果导出"]
tabs = st.tabs(tab_names)


with tabs[0]:
    st.header("使用说明")
    st.markdown("""
    **使用流程**
    1. 侧边栏上传采食、体重 Excel（可多文件多工作表）。
    2. 设置训练结束日期/时间、昼夜时段、异常值参数。
    3. 勾选分析模块与导出项目。
    4. 点击 **开始分析**。
    5. **导出按钮在左侧边栏『⑦ 一键导出』** —— 分析完成后立即可用；也可切到最右边『结果导出』Tab 分类导出。
    6. 各 Tab 顶部有独立绘图参数；修改后图形自动刷新。
    7. 每张图有 **下载PNG** 按钮（需安装 kaleido）。

    **核心指标**

    | 缩写 | 中文名 | 定义 |
    |---|---|---|
    | TFB | 总采食次数 | 有效采食记录总次数 |
    | FI | 总采食量 | 有效采食量之和 (g) |
    | AMS | 平均单次采食量 | FI / TFB (g/bout) |
    | TFD | 总采食时长 | 有效采食时长之和 (s) |
    | AFBD | 平均单次采食时长 | TFD / TFB (s/bout) |
    | IMI | 平均采食间隔 | 相邻两次采食起始时间间隔均值 (s) |
    | FR | 采食速率 | 单次采食量 / 单次采食时长 (g/s) |
    | IBW | 初始体重 | 每只鸭第一条有效体重 (kg) |
    | FBW | 最终体重 | 每只鸭最后一条有效体重 (kg) |
    | Gain | 增重 | FBW − IBW (kg) |
    | ADG | 平均日增重 | Gain / 生长天数 (g/d) |
    | ADFI | 平均日采食量 | FI / 有效采食天数 (g/d) |
    | MBW | 代谢体重 | ((IBW + FBW)/2)^0.75 |
    | FCR | 料重比 | FI / 增重；越小越好 |
    | RFI | 剩余采食量 | 观察 ADFI − 由 MBW 和 ADG 预测的 ADFI |

    **采食时长格式**：`1.09:33:38`、`09:33:38`、`33:33:38`、`09:33` 均可识别。

    **V5.0 更新**
    - 🔧 **导出功能前移到侧边栏**：分析完成后，"⑦ 一键导出"按钮直接出现在侧边栏，不用再翻 Tab
    - Tab 数 14 → 12（合并了"单次采食/生长曲线"、"每周 FCR/相关性"），不再被 Streamlit 折叠
    - 修复分析失败时旧 result 残留的问题
    - 保留 V4.9 的所有修复（Series ambiguous、KeyError N/Mean/SD、xlrd 等）
    """)

    st.info("如果上传后报错，请切换到『QC』标签页查看『读取工作表记录』，"
            "该表会告诉你每个工作表匹配到了哪些字段、缺失哪些字段。"
            "如遇 .xls 读取失败，请先执行：pip install xlrd")


with tabs[1]:
    st.header("数据清洗 QC")

    if "read_qc" in result and len(result["read_qc"]) > 0:
        rq = result["read_qc"]
        failed = rq[rq["状态"].astype(str).str.contains("跳过", na=False)]
        if len(failed) > 0:
            st.warning(
                f"⚠️ 有 {len(failed)} 个工作表被跳过。请看下方『读取工作表记录』的"
                "『状态』和『说明』列：说明会告诉你每个 sheet 缺了哪些字段、"
                "以及 pandas 实际读到的候选列名。"
            )

    if "qc_summary" in result and len(result["qc_summary"]) > 0:
        st.subheader("整体 QC")
        st.dataframe(result["qc_summary"], use_container_width=True)
    if "datetime_qc" in result:
        st.subheader("日期时间解析 QC")
        st.dataframe(result["datetime_qc"], use_container_width=True)
    if "read_qc" in result:
        st.subheader("读取工作表记录")
        st.dataframe(result["read_qc"], use_container_width=True)
    if "week_summary" in result and len(result["week_summary"]) > 0:
        st.subheader("周次数据量检查")
        st.dataframe(result["week_summary"], use_container_width=True)
    if "na_qc" in result:
        st.subheader("重要字段 NA")
        st.dataframe(result["na_qc"], use_container_width=True)
    if "feed_week_stats" in result:
        st.subheader("采食时长周内 QC")
        st.dataframe(result["feed_week_stats"].round(4), use_container_width=True)
    if "bw_week_stats" in result:
        st.subheader("体重周内 QC")
        st.dataframe(result["bw_week_stats"].round(4), use_container_width=True)
    if "feed_outlier" in result:
        st.subheader("异常采食记录")
        st.dataframe(result["feed_outlier"], use_container_width=True)
    if "bw_outlier" in result:
        st.subheader("异常体重记录")
        st.dataframe(result["bw_outlier"], use_container_width=True)


with tabs[2]:
    st.header("每只鸭综合指标")
    if "individual" in result:
        st.dataframe(result["individual"].round(4), use_container_width=True)


with tabs[3]:
    st.header("采食行为（TFB / FI / AMS / TFD / AFBD / IMI / FR）")
    if "feeding" in result:
        cols = ["Animal_ID", "TFB", "FI_g", "AMS_g", "TFD_sec", "AFBD_sec", "IMI_sec", "FR_g_sec"]
        cols = [c for c in cols if c in result["feeding"].columns]
        st.dataframe(result["feeding"][cols].round(4), use_container_width=True)
    if "bout" in result:
        st.subheader("单次采食记录")
        st.dataframe(result["bout"].head(2000), use_container_width=True)


with tabs[4]:
    st.header("生产性能")
    if "production" in result:
        cols = ["Animal_ID", "Group", "IBW_kg", "FBW_kg", "Gain_kg", "ADG_g",
                "ADFI_g", "MBW", "FCR", "RFI", "Feed_Efficiency", "TFB", "FI_g"]
        cols = [c for c in cols if c in result["production"].columns]
        st.dataframe(result["production"][cols].round(4), use_container_width=True)


with tabs[5]:
    st.header("HFF / LFF")
    st.markdown("**本页绘图参数**")
    c1, c2 = st.columns(2)
    with c1:
        hff_palette = st.selectbox("配色方案", list(PALETTES.keys()), index=0, key="hff_pal")
    with c2:
        hff_max_n = st.number_input("每类最大记录数", 500, 50000, 5000, 500, key="hff_n")

    if "hff_group_summary" in result and len(result["hff_group_summary"]) > 0:
        st.subheader("HFF / LFF 分组概况")
        st.dataframe(result["hff_group_summary"].round(2), use_container_width=True)
    if "hff_tests" in result and len(result["hff_tests"]) > 0:
        st.subheader("HFF vs LFF 检验")
        df = result["hff_tests"].copy()
        for c in ["Mann_Whitney_P", "T_test_P"]:
            if c in df.columns:
                df[c] = df[c].apply(format_p)
        st.dataframe(df, use_container_width=True)
    if "hff" in result and len(result["hff"]) > 0:
        st.subheader("HFF / LFF 分组结果")
        st.dataframe(result["hff"].round(4), use_container_width=True)

        plot_dat = result["hff"].dropna(subset=["Feed_Frequency_Group"])
        plot_dat = plot_dat[plot_dat["Feed_Frequency_Group"].isin(["HFF", "LFF"])]
        if len(plot_dat) > 0:
            long = plot_dat.melt(id_vars=["Animal_ID", "Feed_Frequency_Group"],
                                  value_vars=[c for c in ["TFB", "FR_g_sec", "FCR", "RFI"] if c in plot_dat.columns],
                                  var_name="Indicator", value_name="Value").dropna()
            if len(long) > 0:
                pal = seq_colors(hff_palette, 2)
                fig = px.violin(long, x="Feed_Frequency_Group", y="Value",
                                color="Feed_Frequency_Group", facet_col="Indicator",
                                facet_col_wrap=2, box=True, points=False,
                                color_discrete_map={"HFF": pal[0], "LFF": pal[1]})
                fig.update_layout(height=650, showlegend=False)
                st.plotly_chart(fig, use_container_width=True)
                dl_button(fig, "dl_hff", "hff_lff.png")


with tabs[6]:
    st.header("采食节律")
    st.markdown("**本页绘图参数**")
    c1, c2 = st.columns(2)
    with c1:
        rhythm_palette = st.selectbox("配色方案", list(PALETTES.keys()), index=0, key="rhythm_pal")
    with c2:
        rhythm_multi = st.checkbox("按周使用不同颜色", value=True, key="rhythm_multi")

    if "rhythm" in result and len(result["rhythm"]) > 0:
        st.subheader("24 小时采食次数（按周）")
        df = result["rhythm"]
        fig = px.line(df, x="Hour_Block", y="Mean_Bouts_Per_Duck",
                      color="Week_Number", markers=True,
                      color_discrete_sequence=seq_colors(rhythm_palette, df["Week_Number"].nunique()))
        fig.update_layout(height=600, xaxis_title="Hour of Day",
                          yaxis_title="Mean Feeding Bouts / Duck")
        st.plotly_chart(fig, use_container_width=True)
        dl_button(fig, "dl_rhythm", "rhythm.png")

    if "daily_bouts" in result and len(result["daily_bouts"]) > 0:
        st.subheader("每日有效访饲次数")
        fig = px.line(result["daily_bouts"], x="Date", y="Mean_Bouts_Per_Duck", markers=True,
                      color_discrete_sequence=seq_colors(rhythm_palette, 1))
        fig.update_layout(height=500)
        st.plotly_chart(fig, use_container_width=True)
        dl_button(fig, "dl_daily", "daily_bouts.png")

    if "weekly_bouts" in result and len(result["weekly_bouts"]) > 0:
        st.subheader("每周访饲次数")
        fig = px.line(result["weekly_bouts"], x="Date", y="Mean_Bouts_Per_Duck",
                      color="Week_Number", markers=True,
                      color_discrete_sequence=seq_colors(rhythm_palette, result["weekly_bouts"]["Week_Number"].nunique()))
        fig.update_layout(height=500)
        st.plotly_chart(fig, use_container_width=True)
        dl_button(fig, "dl_weekly", "weekly_bouts.png")


with tabs[7]:
    st.header("创新行为指标")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        inno_palette = st.selectbox("散点配色", list(PALETTES.keys()), index=1, key="inno_pal")
    with c2:
        phase_palette = st.selectbox("峰值时刻直方图配色", list(PALETTES.keys()), index=2, key="phase_pal")
    with c3:
        inno_show_se = st.checkbox("显示回归置信带", value=True, key="inno_se")
    with c4:
        inno_line_color = st.text_input("回归线颜色(Hex, 可留空)", "", key="inno_line")

    if "daynight_table" in result:
        st.subheader("① 昼夜采食分配")
        st.dataframe(result["daynight_table"].round(4), use_container_width=True)

    if "innovation_merged" in result and len(result["innovation_merged"]) > 0:
        merged = result["innovation_merged"]
        pal = seq_colors(inno_palette, 3)
        line_col = valid_hex(inno_line_color) if inno_line_color.strip() else pal[1]

        dat = merged.dropna(subset=["Day_FI_Ratio", "FCR"]) if {"Day_FI_Ratio", "FCR"}.issubset(merged.columns) else pd.DataFrame()
        if len(dat) >= 5:
            r, p = stats.spearmanr(dat["Day_FI_Ratio"], dat["FCR"])
            fig = px.scatter(dat, x="Day_FI_Ratio", y="FCR",
                             trendline="ols",
                             trendline_color_override=line_col,
                             color_discrete_sequence=[pal[0]])
            fig.update_traces(marker=dict(size=8, opacity=0.7), selector=dict(mode="markers"))
            fig.update_layout(height=500, title=f"Spearman r={r:.3f}, P={format_p(p)}, n={len(dat)}")
            st.plotly_chart(fig, use_container_width=True)
            dl_button(fig, "dl_dn", "daynight_scatter.png")

        st.subheader("昼夜比 与 生产性能的相关")
        corr = compute_innovation_corr(merged,
                                       ["Day_FI_Ratio", "Day_Bout_Ratio", "Day_Night_FI_Ratio",
                                        "Day_Night_Bout_Ratio", "Day_Avg_Meal", "Night_Avg_Meal"])
        if len(corr) > 0:
            corr["Spearman_r"] = corr["Spearman_r"].round(4)
            corr["P_value"] = corr["P_value"].apply(format_p)
            st.dataframe(corr, use_container_width=True)

    if "cv_table" in result:
        st.subheader("② 行为一致性（CV）")
        st.dataframe(result["cv_table"].round(4), use_container_width=True)

    if "innovation_merged" in result and len(result["innovation_merged"]) > 0:
        merged = result["innovation_merged"]
        cv_vars = ["CV_Duration", "CV_FR", "Robust_CV_IMI", "CV_Daily_Bouts", "CV_Daily_FI", "CV_Daily_TFD"]
        cv_vars = [v for v in cv_vars if v in merged.columns]
        if cv_vars and "FCR" in merged.columns:
            plot_dat = merged.melt(id_vars=["Animal_ID", "FCR"], value_vars=cv_vars,
                                    var_name="CV_Type", value_name="CV_Value").dropna()
            if len(plot_dat) >= 5:
                fig = px.scatter(plot_dat, x="CV_Value", y="FCR", facet_col="CV_Type",
                                 facet_col_wrap=3, trendline="ols",
                                 color_discrete_sequence=[seq_colors(inno_palette, 1)[0]])
                fig.update_layout(height=650)
                st.plotly_chart(fig, use_container_width=True)
                dl_button(fig, "dl_cv", "cv_fcr.png")

        st.subheader("CV 与 生产性能的相关")
        corr = compute_innovation_corr(merged, cv_vars)
        if len(corr) > 0:
            corr["Spearman_r"] = corr["Spearman_r"].round(4)
            corr["P_value"] = corr["P_value"].apply(format_p)
            st.dataframe(corr, use_container_width=True)

    if "cosinor_table" in result:
        st.subheader("③ 昼夜节律余弦拟合")
        st.dataframe(result["cosinor_table"].round(4), use_container_width=True)

        if "innovation_merged" in result and len(result["innovation_merged"]) > 0:
            merged = result["innovation_merged"]
            if {"Cosinor_A", "FCR"}.issubset(merged.columns):
                dat = merged.dropna(subset=["Cosinor_A", "FCR"])
                if len(dat) >= 5:
                    r, p = stats.spearmanr(dat["Cosinor_A"], dat["FCR"])
                    pal = seq_colors(inno_palette, 3)
                    line_col = valid_hex(inno_line_color) if inno_line_color.strip() else pal[1]
                    fig = px.scatter(dat, x="Cosinor_A", y="FCR", trendline="ols",
                                     trendline_color_override=line_col,
                                     color_discrete_sequence=[pal[0]])
                    fig.update_layout(height=500, title=f"Spearman r={r:.3f}, P={format_p(p)}, n={len(dat)}")
                    st.plotly_chart(fig, use_container_width=True)
                    dl_button(fig, "dl_cos", "cosinor_scatter.png")

            if "Peak_Hour" in merged.columns:
                phase_dat = merged.dropna(subset=["Peak_Hour"])
                if len(phase_dat) >= 5:
                    phase_col = seq_colors(phase_palette, 1)[0]
                    fig = px.histogram(phase_dat, x="Peak_Hour", nbins=24,
                                       color_discrete_sequence=[phase_col])
                    fig.update_layout(height=450, xaxis_title="Peak Hour",
                                      yaxis_title="Number of ducks")
                    st.plotly_chart(fig, use_container_width=True)
                    dl_button(fig, "dl_phase", "phase_hist.png")

            st.subheader("余弦参数 与 生产性能的相关")
            corr = compute_innovation_corr(merged, ["Cosinor_A", "Cosinor_R2", "Cosinor_M", "Peak_Hour"])
            if len(corr) > 0:
                corr["Spearman_r"] = corr["Spearman_r"].round(4)
                corr["P_value"] = corr["P_value"].apply(format_p)
                st.dataframe(corr, use_container_width=True)

    if "fano_table" in result:
        st.subheader("④ Fano 聚集指数")
        st.dataframe(result["fano_table"].round(4), use_container_width=True)

        if "innovation_merged" in result and len(result["innovation_merged"]) > 0:
            merged = result["innovation_merged"]
            st.subheader("Fano 与 生产性能的相关")
            corr = compute_innovation_corr(merged, ["Fano", "Mean_Bouts_Per_Hour", "Var_Bouts_Per_Hour"])
            if len(corr) > 0:
                corr["Spearman_r"] = corr["Spearman_r"].round(4)
                corr["P_value"] = corr["P_value"].apply(format_p)
                st.dataframe(corr, use_container_width=True)


# ---------- Tab 9：单次采食 + 生长曲线 ----------
with tabs[8]:
    st.header("单次采食")
    st.markdown("**本页绘图参数**")
    c1, c2, c3 = st.columns(3)
    with c1:
        bout_palette = st.selectbox("配色方案", list(PALETTES.keys()), index=0, key="bout_pal")
    with c2:
        bout_multi = st.checkbox("每周使用不同颜色", value=False, key="bout_multi")
    with c3:
        bout_max_n = st.number_input("每类最大记录数", 500, 50000, 5000, 500, key="bout_n")
    c4, c5 = st.columns(2)
    with c4:
        bout_lq = st.number_input("下限分位数(%)", 0.0, 49.0, 1.0, 0.5, key="bout_lq")
    with c5:
        bout_uq = st.number_input("上限分位数(%)", 51.0, 100.0, 99.0, 0.5, key="bout_uq")

    if "bout" in result and len(result["bout"]) > 0:
        bout = result["bout"]
        pal = seq_colors(bout_palette, max(bout["Week_Number"].nunique(), 1))

        for var, label, fname in [
            ("FR_g_sec", "采食速率 FR", "fr_violin.png"),
            ("IMI_sec", "采食间隔 IMI", "imi_violin.png"),
            ("Feed_Duration_sec", "采食时长", "duration_violin.png"),
        ]:
            if var not in bout.columns:
                continue
            if bout[var].dropna().empty:
                continue
            st.subheader(f"{label}：小提琴图 + 箱线图")
            dat = bout.dropna(subset=[var]).copy()
            dat["Plot_Group"] = "Week " + dat["Week_Number"].astype(int).astype(str)
            lo, hi = np.percentile(dat[var], [bout_lq, bout_uq])
            plot_dat = dat[(dat[var] >= lo) & (dat[var] <= hi)]
            fig = px.violin(plot_dat, x="Plot_Group", y=var, box=True, points=False,
                            color="Plot_Group" if bout_multi else None,
                            color_discrete_sequence=pal if bout_multi else [pal[0]])
            fig.update_layout(height=550, showlegend=False, xaxis_title="Week", yaxis_title=label)
            st.plotly_chart(fig, use_container_width=True)
            dl_button(fig, f"dl_{var}", fname)

    if "bout" in result and "bout_kw" in selected_modules:
        st.subheader("Kruskal-Wallis 检验")
        rows = []
        for var, label in [("FR_g_sec", "Feeding Rate"),
                            ("IMI_sec", "Inter-Meal Interval"),
                            ("Feed_Duration_sec", "Feeding Duration")]:
            dat = result["bout"].dropna(subset=[var, "Week_Number"])
            if len(dat) < 3 or dat["Week_Number"].nunique() < 2:
                rows.append({"指标": label, "Kruskal_Wallis_P": np.nan})
                continue
            groups = [g[var].values for _, g in dat.groupby("Week_Number")]
            try:
                p = stats.kruskal(*groups).pvalue
            except Exception:
                p = np.nan
            rows.append({"指标": label, "Kruskal_Wallis_P": p})
        df = pd.DataFrame(rows)
        df["Kruskal_Wallis_P"] = df["Kruskal_Wallis_P"].apply(format_p)
        st.dataframe(df, use_container_width=True)

    st.divider()
    st.header("生长曲线")
    c1, c2, c3 = st.columns(3)
    with c1:
        growth_palette = st.selectbox("曲线配色", list(PALETTES.keys()), index=0, key="growth_pal")
    with c2:
        growth_multi = st.checkbox("HFF/LFF 用不同颜色", value=True, key="growth_multi")
    with c3:
        growth_overall = st.text_input("总体曲线颜色(Hex, 可留空)", "", key="growth_col")

    if "bw_clean" in result and len(result["bw_clean"]) > 0:
        bw = result["bw_clean"].dropna(subset=["BW_kg", "Experimental_Day"])
        if len(bw) > 0:
            daily = bw.sort_values(["Animal_ID", "Date", "Time1"]).groupby(
                ["Animal_ID", "Date", "Experimental_Day"]).tail(1)
            overall = daily.groupby("Experimental_Day").agg(
                Mean_BW_kg=("BW_kg", "mean"),
                SD_BW_kg=("BW_kg", "std"),
                N=("BW_kg", "size"),
            ).reset_index().sort_values("Experimental_Day")

            if len(overall) > 0:
                st.subheader("总体生长曲线")
                pal = seq_colors(growth_palette, 1)
                col = valid_hex(growth_overall) if growth_overall.strip() else pal[0]
                fig = px.line(overall, x="Experimental_Day", y="Mean_BW_kg", markers=True,
                              color_discrete_sequence=[col])
                fig.update_layout(height=550, xaxis_title="Experimental Day",
                                  yaxis_title="Daily Terminal Body Weight (kg)")
                st.plotly_chart(fig, use_container_width=True)
                dl_button(fig, "dl_growth", "growth_overall.png")

            if "production" in result and "Feed_Frequency_Group" in result["production"].columns:
                st.subheader("HFF / LFF 分组生长曲线")
                grp_map = result["production"][["Animal_ID", "Feed_Frequency_Group"]]
                gdat = daily.merge(grp_map, on="Animal_ID", how="left")
                gdat = gdat[gdat["Feed_Frequency_Group"].isin(["HFF", "LFF"])]
                gsum = gdat.groupby(["Feed_Frequency_Group", "Experimental_Day"]).agg(
                    Mean_BW_kg=("BW_kg", "mean")).reset_index().sort_values(["Feed_Frequency_Group", "Experimental_Day"])
                if len(gsum) > 0:
                    pal2 = seq_colors(growth_palette, 2)
                    fig = px.line(gsum, x="Experimental_Day", y="Mean_BW_kg",
                                  color="Feed_Frequency_Group", markers=True,
                                  color_discrete_map={"HFF": pal2[0], "LFF": pal2[1]})
                    fig.update_layout(height=550, xaxis_title="Experimental Day",
                                      yaxis_title="Daily Terminal Body Weight (kg)")
                    st.plotly_chart(fig, use_container_width=True)
                    dl_button(fig, "dl_ggrowth", "growth_by_group.png")


# ---------- Tab 10：每周 FCR + 相关性 ----------
with tabs[9]:
    st.header("每周 FCR")
    c1, c2, c3 = st.columns(3)
    with c1:
        wfcr_palette = st.selectbox("配色方案", list(PALETTES.keys()), index=0, key="wfcr_pal")
    with c2:
        wfcr_multi = st.checkbox("每周使用不同颜色", value=False, key="wfcr_multi")
    with c3:
        wfcr_max_n = st.number_input("每类最大记录数", 500, 50000, 5000, 500, key="wfcr_n")
    c4, c5 = st.columns(2)
    with c4:
        wfcr_lq = st.number_input("下限分位数(%)", 0.0, 49.0, 1.0, 0.5, key="wfcr_lq")
    with c5:
        wfcr_uq = st.number_input("上限分位数(%)", 51.0, 100.0, 99.0, 0.5, key="wfcr_uq")

    if "weekly_fcr" in result and len(result["weekly_fcr"]) > 0:
        dat = result["weekly_fcr"].dropna(subset=["Weekly_FCR"]).copy()
        if len(dat) > 0:
            dat["Week"] = "Week " + dat["Week_Number"].astype(int).astype(str)
            lo, hi = np.percentile(dat["Weekly_FCR"], [wfcr_lq, wfcr_uq])
            plot_dat = dat[(dat["Weekly_FCR"] >= lo) & (dat["Weekly_FCR"] <= hi)]
            pal = seq_colors(wfcr_palette, dat["Week"].nunique())
            fig = px.violin(plot_dat, x="Week", y="Weekly_FCR", box=True, points=False,
                            color="Week" if wfcr_multi else None,
                            color_discrete_sequence=pal if wfcr_multi else [pal[0]])
            fig.update_layout(height=600, showlegend=False)
            st.plotly_chart(fig, use_container_width=True)
            dl_button(fig, "dl_wfcr", "weekly_fcr.png")

    if "weekly_fcr_tests" in result:
        st.subheader("Kruskal-Wallis 检验")
        df = result["weekly_fcr_tests"].copy()
        df["P值"] = df["P值"].apply(format_p)
        st.dataframe(df, use_container_width=True)

    st.divider()
    st.header("12指标 Spearman 相关矩阵")
    corr_palette = st.selectbox("热图配色", list(HEAT_PALETTES.keys()), index=0, key="corr_pal")

    if "correlation" in result and result["correlation"] is not None and len(result["correlation"]) > 0:
        mat = result["correlation"]
        heat_cols = HEAT_PALETTES[corr_palette]
        fig = px.imshow(mat.values, x=mat.columns, y=mat.index,
                        color_continuous_scale=[heat_cols[0], heat_cols[1], heat_cols[2]],
                        zmin=-1, zmax=1, text_auto=".2f", aspect="auto")
        fig.update_layout(height=750)
        st.plotly_chart(fig, use_container_width=True)
        dl_button(fig, "dl_corr", "correlation.png")

        if "correlation_p" in result:
            st.subheader("P 值矩阵")
            st.dataframe(result["correlation_p"].round(4), use_container_width=True)


with tabs[10]:
    st.header("清洗后数据")
    if "feed_clean" in result:
        st.subheader("采食数据")
        st.dataframe(result["feed_clean"].head(2000), use_container_width=True)
    if "bw_clean" in result:
        st.subheader("体重数据")
        st.dataframe(result["bw_clean"].head(2000), use_container_width=True)
    if "bout" in result:
        st.subheader("完整单次采食数据")
        st.dataframe(result["bout"].head(2000), use_container_width=True)


# ============================================================
# 结果导出（V5.0，侧边栏也有一键导出）
# ============================================================
with tabs[11]:
    st.header("Excel 结果导出")

    if not result:
        st.info("请先在左侧点击 **开始分析**，分析完成后本页会显示导出选项。"
                "（也可以直接使用侧边栏『⑦ 一键导出』）")
    else:
        st.success("✅ 分析已完成。下方可预览导出内容，并进行一键导出 / 分类导出 / CSV 下载。")
        st.caption("提示：侧边栏『⑦ 一键导出』按钮同样可用，无需停留在本页。")

        st.subheader("① 导出内容预览")
        preview_rows = []
        for key in selected_exports:
            for name, df in _collect_export_sheets(result, key):
                if df is not None and len(df) > 0:
                    preview_rows.append({
                        "导出类别": EXPORT_MAP_KEYS.get(key, key),
                        "Sheet 名称": name,
                        "行数": len(df),
                        "列数": len(df.columns),
                        "状态": "✅ 有数据",
                    })
                else:
                    preview_rows.append({
                        "导出类别": EXPORT_MAP_KEYS.get(key, key),
                        "Sheet 名称": name,
                        "行数": 0,
                        "列数": 0,
                        "状态": "⚠️ 空（该模块未运行或被跳过）",
                    })
        if preview_rows:
            st.dataframe(pd.DataFrame(preview_rows), use_container_width=True)
        else:
            st.warning("你还没有选择任何导出项，请到左侧『⑥ Excel 导出项目』勾选。")

        st.subheader("② 一键导出（合并成一个 Excel）")
        if not selected_exports:
            st.warning("请在左侧『⑥ Excel 导出项目』勾选至少一个导出项。")
        else:
            ts = datetime.today().strftime("%Y%m%d_%H%M%S")
            try:
                xls_bytes = build_excel(result, selected_exports, params_info=_params_info())
                st.download_button(
                    label=f"📥 导出全部所选（{len(selected_exports)} 个类别）",
                    data=xls_bytes,
                    file_name=f"肉鸭分析_V5.0_{ts}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key="page_dl_all",
                )
                st.caption(f"文件大小：约 {len(xls_bytes)/1024:.1f} KB")
            except Exception as e:
                st.error(f"合并导出失败：{e}")
                import traceback
                st.code(traceback.format_exc())

        st.subheader("③ 按类别单独导出")
        st.caption("每个类别生成一个独立的 Excel 文件，便于分发给不同的人。")
        ts2 = datetime.today().strftime("%Y%m%d_%H%M%S")
        for key in selected_exports:
            label = EXPORT_MAP_KEYS.get(key, key)
            try:
                xls_bytes = build_excel(result, [key], params_info=_params_info())
                st.download_button(
                    label=f"📥 {label}",
                    data=xls_bytes,
                    file_name=f"肉鸭_{key}_{ts2}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key=f"page_dl_{key}",
                )
            except Exception as e:
                st.error(f"「{label}」导出失败：{e}")

        with st.expander("④ 附加：核心数据单独下载 CSV（便于快速查看）"):
            csv_items = [
                ("个体综合指标", result.get("individual")),
                ("采食行为", result.get("feeding")),
                ("生产性能", result.get("production")),
                ("清洗后采食数据", result.get("feed_clean")),
                ("清洗后体重数据", result.get("bw_clean")),
            ]
            for name, df in csv_items:
                if df is not None and len(df) > 0:
                    csv_bytes = df.to_csv(index=False).encode("utf-8-sig")
                    st.download_button(
                        f"⬇️ {name}.csv",
                        data=csv_bytes,
                        file_name=f"{name}_{ts2}.csv",
                        mime="text/csv",
                        key=f"csv_{name}",
                    )